# -*- coding: utf-8 -*-
"""结果报表（xlsx）。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .publish_settings import format_publish_at, summarize_author

HEADERS = [
    ("账号", 14),
    ("序号", 6),
    ("文件名", 34),
    ("广告语", 40),
    ("作者声明", 24),
    ("发布时间", 18),
    ("发布方式", 12),
    ("发布状态", 14),
    ("作品链接", 40),
    ("金牛改名", 12),
    ("跳过原因", 30),
    ("待改名", 10),
    ("耗时(秒)", 10),
    ("更新时间", 20),
    ("备注", 30),
]


def write_report(path: Path, items: List[Dict[str, Any]]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "运行结果"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4F81BD")
    for col, (title, width) in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "A2"

    def sort_key(item: Dict[str, Any]):
        return (str(item.get("account", "")), int(item.get("index", 0)))

    for row, item in enumerate(sorted(items, key=sort_key), start=2):
        publish_at = item.get("publish_at", "")
        publish_mode = item.get("publish_mode") or ("定时发布" if publish_at else "立即发布")
        rename_status = item.get("rename_status", "")
        pending = "是" if publish_mode == "定时发布" and rename_status not in ("已改名",) else ""
        values = [
            item.get("account", ""),
            int(item.get("index", 0)) + 1,
            item.get("file", ""),
            item.get("copy", ""),
            summarize_author(item.get("_author_statement", "")),
            format_publish_at(publish_at),
            publish_mode,
            item.get("publish_status", ""),
            item.get("work_url", ""),
            rename_status,
            item.get("skip_reason", ""),
            pending,
            item.get("seconds", 0),
            item.get("updated", ""),
            item.get("note", ""),
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=(col in (3, 4, 5, 11, 15)))

    summary = book.create_sheet("汇总")
    summary.append(["账号", "总条数", "已发布", "定时发布", "发布异常", "已改名", "待改名"])
    for cell in summary[1]:
        cell.font = head_font
        cell.fill = head_fill
    accounts: Dict[str, List[Dict[str, Any]]] = {}
    for item in items:
        accounts.setdefault(str(item.get("account", "")), []).append(item)
    for name, group in sorted(accounts.items()):
        summary.append(
            [
                name,
                len(group),
                sum(1 for it in group if it.get("publish_status") == "已发布"),
                sum(1 for it in group if it.get("publish_mode") == "定时发布"),
                sum(1 for it in group if it.get("publish_status") in ("失败", "发布结果不确定")),
                sum(1 for it in group if it.get("rename_status") == "已改名"),
                sum(
                    1
                    for it in group
                    if it.get("publish_mode") == "定时发布" and it.get("rename_status") != "已改名"
                ),
            ]
        )
    for col, width in enumerate((16, 10, 10, 12, 12, 10, 10), start=1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary.cell(row=len(accounts) + 3, column=1, value="生成时间")
    summary.cell(row=len(accounts) + 3, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    book.save(path)
    return path
